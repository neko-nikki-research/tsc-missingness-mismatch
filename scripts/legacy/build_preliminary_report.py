"""Build a concise Word report from the completed benchmark outputs."""

from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
RESULTS = ROOT / "results" / "rate_expanded"
OUTPUT = ROOT / "preliminary_experiment_report.docx"


def shade(cell, color: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    fill = OxmlElement("w:shd")
    fill.set(qn("w:fill"), color)
    properties.append(fill)


def set_cell_text(cell, text: str, bold: bool = False, color: str | None = None) -> None:
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(9)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(document: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        shade(table.rows[0].cells[index], "1F4E78")
        set_cell_text(table.rows[0].cells[index], header, bold=True, color="FFFFFF")
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(row):
            if row_index % 2 == 1:
                shade(cells[index], "EAF2F8")
            set_cell_text(cells[index], value)


def main() -> None:
    summary = pd.read_csv(RESULTS / "tables" / "regret_by_rate_and_condition.csv")
    test = pd.read_csv(RESULTS / "tables" / "paired_wilcoxon.csv").iloc[0]
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    styles = document.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10.5)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("时间序列分类缺失模式不匹配 初步实验报告")
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("UCR 数据集上的验证阶段与部署阶段缺失机制比较").italic = True

    document.add_heading("主要结论", level=1)
    document.add_paragraph(
        "本项目考察验证阶段和部署阶段的时间序列缺失模式不一致时，基于验证集选择的分类器是否会在部署数据上产生性能损失。"
        "在当前 3 个 UCR 数据集、5 个随机种子和 10% 至 30% 缺失率的实验中，不匹配条件的平均 selection regret 高于匹配条件。"
        "这一差异在 30% 缺失率时最明显。但配对 Wilcoxon 检验的单侧 p 值为 "
        f"{test['one_sided_p_value']:.3f}，尚未达到常用的统计显著性标准，因此本报告将其解释为需要更多数据集验证的趋势。"
    )

    document.add_heading("研究问题", level=1)
    document.add_paragraph(
        "How does mismatch between validation-time and deployment-time temporal missingness patterns affect classifier selection and deployment performance in time-series classification?"
    )

    document.add_heading("实验设计", level=1)
    add_table(document, ["项目", "设置"], [
        ["数据集", "GunPoint、ECG200、ItalyPowerDemand"],
        ["数据划分", "官方 train 分层拆分为 train 和 validation；官方 test 仅模拟部署"],
        ["缺失率", "10%、20%、30%"],
        ["模式", "point 与 block；matched 为相同模式，mismatched 为不同模式"],
        ["填补", "线性插值 linear interpolation"],
        ["分类器", "1NN-DTW、Time Series Forest、MiniROCKET"],
        ["重复次数", "3 数据集 × 5 seeds × 3 缺失率 × 4 模式组合 × 3 分类器"],
        ["选择规则", "只依据 validation accuracy 选择；test 只用于 oracle 与 regret"],
    ])

    document.add_heading("指标定义", level=1)
    document.add_paragraph("Selected model：validation accuracy 最高的分类器。")
    document.add_paragraph("Oracle model：test accuracy 最高的分类器，仅用于事后比较。")
    document.add_paragraph("Selection regret：oracle test accuracy 减去 selected model test accuracy。")

    document.add_heading("结果", level=1)
    rows = []
    for rate in sorted(summary.missing_rate.unique()):
        subset = summary[summary.missing_rate == rate].set_index("condition")
        rows.append([
            f"{rate * 100:.0f}%",
            f"{subset.loc['matched', 'mean_regret']:.4f}",
            f"{subset.loc['mismatched', 'mean_regret']:.4f}",
            f"{subset.loc['matched', 'selection_error_rate'] * 100:.1f}%",
            f"{subset.loc['mismatched', 'selection_error_rate'] * 100:.1f}%",
        ])
    add_table(document, ["缺失率", "Matched regret", "Mismatched regret", "Matched 选错率", "Mismatched 选错率"], rows)
    document.add_paragraph(
        "在 10% 缺失率下，两种条件的 regret 接近；20% 时 mismatch 略高；30% 时 mismatch 的平均 regret 为 0.0293，"
        "高于 matched 的 0.0111。误差条反映不同数据集与随机种子之间的波动。"
    )

    figure = document.add_paragraph()
    figure.alignment = WD_ALIGN_PARAGRAPH.CENTER
    figure.add_run().add_picture(str(RESULTS / "figures" / "regret_by_rate.png"), width=Inches(6.1))
    caption = document.add_paragraph("图 1  不同缺失率下 matched 与 mismatched 的 selection regret")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.runs[0].italic = True

    document.add_heading("统计检验与限制", level=1)
    document.add_paragraph(
        f"对同一数据集、seed、缺失率和 source pattern 的 matched 与 mismatched regret 进行配对 Wilcoxon signed-rank test。"
        f"共有 {int(test['n_pairs'])} 对；mismatch 减 matched 的平均差为 {test['mean_difference_mismatch_minus_matched']:.4f}，"
        f"单侧 p 值为 {test['one_sided_p_value']:.3f}。当前样本量和数据集数量有限，不能据此断言总体显著差异。"
    )
    document.add_paragraph(
        "本阶段只使用 linear interpolation，且正式缺失率实验只比较 point 和 block。后续应加入更多 UCR 数据集、更多 seed、"
        "zero 和 mean 填补，以及 prefix 和 suffix 缺失，以检验结论的稳健性。"
    )

    document.add_heading("可复现性", level=1)
    document.add_paragraph(
        "正式缺失率实验配置为 configs/rate_expanded.yaml；原始分类器记录为 results/rate_expanded/raw_results.csv；"
        "模型选择记录为 results/rate_expanded/selection_results.csv。所有 masking 使用固定 seed，且所有分类器在每次比较中共享完全相同的 mask。"
    )
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
